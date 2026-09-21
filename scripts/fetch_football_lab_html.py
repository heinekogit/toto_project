#!/usr/bin/env python3
import argparse
import csv
import html
import os
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import List
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import load_workbook

from http_retry import get_with_retry


ROOT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT_DIR / "data" / "manual" / "lab情報一覧.xlsx"
DEFAULT_OUT_ROOT = ROOT_DIR / "external_contents"
USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/136.0.0.0 Safari/537.36"
)


def latest_prediction_snapshot(league: str, season_year: int) -> Path | None:
    league_key = _league_key(league)
    snapshot_dir = ROOT_DIR / "data" / "output_snapshots" / f"{league_key}_{season_year}"
    pattern = f"*_predictions_candidate_{league_key}_{season_year}_predictions_hfa_on.csv"
    candidates = [path for path in snapshot_dir.glob(pattern) if path.is_file()]
    return max(candidates, key=lambda path: path.name) if candidates else None


def print_snapshot_commands(out_root: Path, snapshot_date: str, leagues: List[str], season_year: int) -> None:
    print("[NEXT] Football LAB snapshot生成コマンド:")
    for league in leagues:
        prediction_csv = latest_prediction_snapshot(league, season_year)
        if prediction_csv is None:
            print(f"[WARN] {league.upper()} の予測snapshotが見つかりません")
            continue
        league_dir = out_root / snapshot_date / league
        try:
            league_dir_text = str(league_dir.relative_to(ROOT_DIR))
            prediction_text = str(prediction_csv.relative_to(ROOT_DIR))
        except ValueError:
            league_dir_text = str(league_dir)
            prediction_text = str(prediction_csv)
        print(
            "./scripts/.venv/bin/python scripts/build_football_lab_snapshot.py \\\n"
            f"  --league-dir {league_dir_text} \\\n"
            f"  --prediction-csv {prediction_text}"
        )


@dataclass
class Target:
    league: str
    data_name: str
    item_name: str
    url: str


def _safe_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return text


def _norm_text(value: object) -> str:
    return unicodedata.normalize("NFKC", _safe_text(value))


def _sanitize_filename(name: str) -> str:
    text = _norm_text(name)
    text = re.sub(r'[\\/:*?"<>|]+', "_", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text or "unnamed"


def _league_key(value: str) -> str:
    text = _norm_text(value).lower()
    if text in {"j1", "j2", "j3"}:
        return text
    raise ValueError(f"unknown league label: {value}")


def load_targets(xlsx_path: Path) -> List[Target]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    targets: List[Target] = []
    league = ""
    data_name = ""
    for row in ws.iter_rows(values_only=True):
        cells = list(row[:4]) + [None] * max(0, 4 - len(row))
        col_league, col_data_name, col_item_name, col_url = cells[:4]
        text_league = _safe_text(col_league)
        text_data_name = _safe_text(col_data_name)
        item_name = _safe_text(col_item_name)
        url = _safe_text(col_url)
        if text_league:
            league = _league_key(text_league)
        if text_data_name:
            data_name = _norm_text(text_data_name)
        if not item_name or not url or item_name == "項目名" or url == "url":
            continue
        if not league:
            raise RuntimeError(f"league missing before row item={item_name} url={url}")
        if not data_name:
            raise RuntimeError(f"data_name missing before row item={item_name} url={url}")
        targets.append(
            Target(
                league=league,
                data_name=data_name,
                item_name=_norm_text(item_name),
                url=url,
            )
        )
    return targets


def write_index_csv(path: Path, rows: List[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "snapshot_date",
        "league",
        "data_name",
        "item_name",
        "url",
        "saved_path",
        "status",
        "http_status",
        "error",
    ]
    merged = {}
    if path.exists():
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                key = (row.get("league", ""), row.get("data_name", ""), row.get("item_name", ""))
                merged[key] = row
    for row in rows:
        key = (row.get("league", ""), row.get("data_name", ""), row.get("item_name", ""))
        merged[key] = row
    ordered_rows = sorted(
        merged.values(),
        key=lambda row: (row.get("league", ""), row.get("data_name", ""), row.get("item_name", "")),
    )
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(ordered_rows)


def build_saved_from_comment(url: str) -> str:
    return f"<!-- saved from url=({len(url):04d}){url} -->\n"


def football_lab_season_label(season_year: int) -> str:
    return f"{int(season_year)}/{str(int(season_year) + 1)[-2:]}"


def build_season_url(url: str, league: str, season_year: int) -> str:
    """Convert legacy special-season URLs to the regular-season URL scheme."""
    league_key = _league_key(league)
    parts = urlsplit(url)
    path = re.sub(r"/j100[12](?=/|$)", f"/{league_key}", parts.path)
    query = dict(parse_qsl(parts.query, keep_blank_values=True))
    query["year"] = str(int(season_year))
    return urlunsplit((parts.scheme, parts.netloc, path, urlencode(query), parts.fragment))


def extract_title(html_text: str) -> str:
    match = re.search(r"<title[^>]*>(.*?)</title>", html_text, flags=re.I | re.S)
    if not match:
        return ""
    text = re.sub(r"<[^>]+>", "", match.group(1))
    return unicodedata.normalize("NFKC", html.unescape(text)).strip()


def validate_season_html(html_text: str, season_year: int, url: str) -> str:
    title = extract_title(html_text)
    expected = football_lab_season_label(season_year)
    if expected not in title:
        raise RuntimeError(
            f"Football LAB season mismatch: expected={expected} title={title!r} url={url}"
        )
    return title


def fetch_and_save(
    target: Target,
    out_root: Path,
    snapshot_date: str,
    dry_run: bool,
    season_year: int,
) -> dict:
    league_dir = out_root / snapshot_date / target.league
    file_name = f"{_sanitize_filename(target.item_name)}.html"
    saved_path = league_dir / file_name
    effective_url = build_season_url(target.url, target.league, season_year)
    row = {
        "snapshot_date": snapshot_date,
        "league": target.league,
        "data_name": target.data_name,
        "item_name": target.item_name,
        "url": effective_url,
        "saved_path": str(saved_path),
        "status": "pending",
        "http_status": "",
        "error": "",
    }
    if dry_run:
        row["status"] = "dry_run"
        return row

    headers = {"User-Agent": USER_AGENT}
    try:
        response = get_with_retry(effective_url, headers=headers, timeout=(5, 30), max_retries=3)
        html_text = response.text
        validate_season_html(html_text, season_year, effective_url)
        if "saved from url=" not in html_text[:300]:
            html_text = build_saved_from_comment(effective_url) + html_text
        league_dir.mkdir(parents=True, exist_ok=True)
        saved_path.write_text(html_text, encoding="utf-8")
        row["status"] = "ok"
        row["http_status"] = str(response.status_code)
        return row
    except Exception as e:
        row["status"] = "error"
        row["error"] = repr(e)
        return row


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="lab情報一覧.xlsx を読んで Football LAB HTML を保存する")
    p.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="既定: data/manual/lab情報一覧.xlsx")
    p.add_argument("--out-root", default=str(DEFAULT_OUT_ROOT), help="既定: external_contents")
    p.add_argument("--date", default=datetime.now().strftime("%Y%m%d"), help="保存日ディレクトリ。既定: 今日")
    p.add_argument(
        "--season-year",
        type=int,
        default=int(os.environ.get("FOOTBALL_LAB_SEASON_YEAR", "2026")),
        help="Football LABのyear。2026は2026/27シーズン（既定: 2026）",
    )
    p.add_argument("--league", choices=["all", "j1", "j2"], default="all", help="取得対象リーグ")
    p.add_argument("--data-name", default="", help="データ名で絞る。例: チームデータ")
    p.add_argument("--dry-run", action="store_true", help="取得せず対象一覧だけ出す")
    return p.parse_args()


def main() -> None:
    args = parse_args()
    xlsx_path = Path(args.xlsx).resolve()
    out_root = Path(args.out_root).resolve()
    snapshot_date = str(args.date).strip()

    targets = load_targets(xlsx_path)
    if args.league != "all":
        targets = [t for t in targets if t.league == args.league]
    if args.data_name:
        data_name_filter = _norm_text(args.data_name)
        targets = [t for t in targets if t.data_name == data_name_filter]

    if not targets:
        raise RuntimeError("取得対象がありません")

    results = [
        fetch_and_save(t, out_root, snapshot_date, args.dry_run, args.season_year)
        for t in targets
    ]
    index_path = out_root / snapshot_date / "football_lab_fetch_index.csv"
    write_index_csv(index_path, results)

    ok_count = sum(1 for r in results if r["status"] == "ok")
    err_count = sum(1 for r in results if r["status"] == "error")
    dry_count = sum(1 for r in results if r["status"] == "dry_run")

    print(f"[OK] index={index_path}")
    print(
        f"[INFO] targets={len(results)} ok={ok_count} error={err_count} dry_run={dry_count} "
        f"league={args.league} data_name={args.data_name or 'all'} snapshot_date={snapshot_date} "
        f"season_year={args.season_year} season_label={football_lab_season_label(args.season_year)}"
    )
    if err_count:
        for row in results:
            if row["status"] == "error":
                print(f"[ERROR] {row['league']} {row['item_name']} {row['url']} :: {row['error']}")
        raise RuntimeError(
            f"Football LAB取得が未完了です: error={err_count}/{len(results)} "
            "（indexを確認し、不完全なリーグはsnapshot化しないでください）"
        )
    if not args.dry_run:
        leagues = sorted({target.league for target in targets})
        print_snapshot_commands(out_root, snapshot_date, leagues, args.season_year)


if __name__ == "__main__":
    main()
